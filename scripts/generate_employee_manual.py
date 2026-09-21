from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "USER_GUIDE.md"
OUTPUT = ROOT / "docs" / "Huong-Dan-Su-Dung-Coffee-Garden.docx"

BLACK = "000000"
DARK_GREEN = "244C3A"
PALE_GREEN = "EEF4F0"
PALE_GRAY = "F5F6F5"
MID_GRAY = "66706B"
LIGHT_GRAY = "D9D9D9"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), LIGHT_GRAY)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name: str = "Aptos") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)


def suppress_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def add_inline(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Cascadia Mono"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(DARK_GREEN)
            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Cascadia Mono")
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Cascadia Mono")
        set_run_font(run, run.font.name or "Aptos")
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run)


def configure_document(doc: Document) -> None:
    doc.core_properties.title = "Hướng dẫn sử dụng Coffee Garden"
    doc.core_properties.subject = "Hướng dẫn vận hành cho nhân viên, giám sát và chủ quán Coffee Garden"
    doc.core_properties.author = "Coffee Garden"

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(32)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.paragraph_format.space_after = Pt(14)

    for style_name, size, before, after in (
        ("Heading 1", 22, 18, 9),
        ("Heading 2", 16, 14, 6),
        ("Heading 3", 12.5, 10, 4),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for list_style in ("List Bullet", "List Number"):
        style = styles[list_style]
        style.font.name = "Aptos"
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_after = Pt(3)


def add_footer(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Hướng dẫn sử dụng Coffee Garden  •  Trang ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    set_run_font(run)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_cover(doc: Document) -> None:
    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_before = Pt(52)
    label.paragraph_format.space_after = Pt(18)
    run = label.add_run("VẬN HÀNH COFFEE GARDEN")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(DARK_GREEN)
    run.font.letter_spacing = Pt(1.2)
    set_run_font(run)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    suppress_paragraph_borders(title)
    title.add_run("Hướng dẫn sử dụng cho nhân viên và quản lý")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    run = subtitle.add_run("Gọi món, bếp, thanh toán, báo cáo, kho, phân quyền nhân viên và chốt ngày")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    set_run_font(run)

    details = (
        ("Đối tượng", "Nhân viên phục vụ, bếp, giám sát, kế toán, quản lý và chủ quán"),
        ("Mục đích", "Cung cấp tài liệu thực hành cho mọi trang vận hành và quy trình hằng ngày"),
        ("Phiên bản", "Bản review tại local  •  20 tháng 9 năm 2026"),
    )
    for key, value in details:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(7)
        label = paragraph.add_run(f"{key}  ")
        label.bold = True
        set_run_font(label)
        add_inline(paragraph, value)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(28)
    run = note.add_run("Dùng tài khoản do quán cấp. Luôn đăng xuất khi rời thiết bị dùng chung.")
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string(DARK_GREEN)
    set_run_font(run)

    doc.add_page_break()


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total_width = 17.1
    if column_count == 2:
        widths = [4.2, 12.9]
    elif column_count == 3:
        widths = [4.2, 4.4, 8.5]
    elif column_count == 4:
        widths = [3.3, 4.0, 3.0, 6.8]
    else:
        widths = [total_width / column_count] * column_count

    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        keep_row_together(row)
        if row_index == 0:
            repeat_header(row)
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.width = Cm(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, DARK_GREEN if row_index == 0 else (PALE_GRAY if row_index % 2 == 0 else WHITE))
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(9)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(WHITE)
    set_table_borders(table)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)


PAGE_BREAK_HEADINGS = {
    "Danh mục trang",
    "Tạo order giúp khách",
    "Xử lý order tại bếp",
    "Xem dashboard và doanh thu",
    "Quản lý nguyên liệu, công thức và tồn kho",
    "Tạo tài khoản và phân quyền nhân viên",
    "Hoàn tất chốt ngày",
    "Các lỗi thường gặp",
}


def parse_guide(doc: Document) -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_lines
        if paragraph_lines:
            paragraph = doc.add_paragraph()
            add_inline(paragraph, " ".join(part.strip() for part in paragraph_lines))
            paragraph_lines = []

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        if stripped.startswith("# "):
            flush_paragraph()
            index += 1
            continue
        if stripped.startswith("## "):
            flush_paragraph()
            heading = stripped[3:].strip()
            if heading in PAGE_BREAK_HEADINGS and len(doc.paragraphs) > 3:
                doc.add_page_break()
            paragraph = doc.add_paragraph(style="Heading 1")
            paragraph.add_run(heading)
            index += 1
            continue
        if stripped.startswith("### "):
            flush_paragraph()
            paragraph = doc.add_paragraph(style="Heading 2")
            paragraph.add_run(stripped[4:].strip())
            index += 1
            continue
        if stripped.startswith("| "):
            flush_paragraph()
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows: list[list[str]] = []
            for table_index, table_line in enumerate(table_lines):
                values = [value.strip() for value in table_line.strip("|").split("|")]
                if table_index == 1 and all(re.fullmatch(r":?-{3,}:?", value) for value in values):
                    continue
                rows.append(values)
            add_table(doc, rows)
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if numbered:
            flush_paragraph()
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.65)
            paragraph.paragraph_format.first_line_indent = Cm(-0.55)
            paragraph.paragraph_format.space_after = Pt(3)
            number_run = paragraph.add_run(f"{numbered.group(1)}.  ")
            set_run_font(number_run)
            add_inline(paragraph, numbered.group(2))
            index += 1
            continue
        if stripped.startswith("- "):
            flush_paragraph()
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
            index += 1
            continue
        paragraph_lines.append(stripped)
        index += 1
    flush_paragraph()


def build() -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    parse_guide(doc)
    add_footer(doc.sections[0])

    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True
        if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}:
            paragraph.paragraph_format.keep_with_next = True

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
